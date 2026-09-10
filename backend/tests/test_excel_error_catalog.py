"""`.xls` と `.xlsx` の両方を読めること（Issue #106）。

**Fixtures are synthetic, and built here rather than committed.** The real
添削資料 may not be copied into this repository in any form
(``AGENTS.md`` "Security"), and a committed binary would in any case be a
blob nobody can review. So both workbooks are written from invented strings
inside the test.

The ``.xlsx`` side is written with openpyxl -- the same library that reads it,
which is fair because what is under test is this module's grid mapping, not
openpyxl's own round trip. The ``.xls`` side has no writer available
(``xlrd`` 2.x reads only, and adding a legacy writer as a third dependency for
one fixture is not worth it), so :func:`_biff2_workbook` emits a minimal
BIFF2 stream by hand -- ~15 lines of ``struct``, no dependency, and its
strings survive as text.

**What that fixture does and does not prove.** It proves that a ``.xls``
suffix is routed to ``xlrd`` and that whatever ``xlrd`` returns is mapped into
the grid the domain reads. It does **not** prove that ``xlrd`` reads BIFF8,
which is what the 6 real ``.xls`` files are -- that is xlrd's own contract,
and the evidence for it is the real-material count recorded in
docs/grading-material-structure.md section 5.4, not this test.
"""

from __future__ import annotations

import logging
import struct
from collections.abc import Callable, Sequence
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import pytest

from auto_scoring.adapters.excel_error_catalog import (
    UnsupportedCatalogFormat,
    annotation_resource_catalog,
    read_error_catalog,
)
from auto_scoring.adapters.local_storage import LocalFileStore
from auto_scoring.domain.error_catalog import ErrorCatalogUnreadable
from auto_scoring.domain.intake_template import MaterialRole
from auto_scoring.domain.test_material import TestMaterial

_HEADER = ["回数", "問題番号", "場所", "生徒の誤り方・現状", "⇒", "採点基準", "赤入れ案"]
_TITLE = ["架空の講座名"]
_DATA = ["第1回", "問1", "", "誤答A", "", "3点減", "赤入れA"]


#: A fixture writer. Rows carry ``object`` rather than ``str`` because a real
#: sheet stores 問題番号 as a number, and that is a case worth writing.
_WriteWorkbook = Callable[[Path, Sequence[Sequence[object]]], Path]


def _write_xlsx(path: Path, rows: Sequence[Sequence[object]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


def _biff2_workbook(path: Path, rows: Sequence[Sequence[object]]) -> Path:
    """A minimal single-sheet BIFF2 ``.xls`` stream: BOF, CODEPAGE, LABEL per
    cell, EOF. Enough for ``xlrd.open_workbook`` to return a sheet."""
    out = bytearray()

    def record(code: int, payload: bytes) -> None:
        out.extend(struct.pack("<HH", code, len(payload)) + payload)

    record(0x0009, struct.pack("<HH", 0x0002, 0x0010))  # BOF, BIFF2 worksheet
    record(0x0042, struct.pack("<H", 932))  # CODEPAGE: the encoding of the LABELs below
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            encoded = str(value).encode("cp932")
            # LABEL: row, column, 3 cell-attribute bytes, length, bytes.
            record(
                0x0004, struct.pack("<HH", r, c) + b"\x00\x00\x00" + bytes([len(encoded)]) + encoded
            )
    record(0x000A, b"")  # EOF
    path.write_bytes(bytes(out))
    return path


@pytest.mark.parametrize(
    ("write", "suffix"), [(_write_xlsx, ".xlsx"), (_biff2_workbook, ".xls")], ids=["xlsx", "xls"]
)
def test_both_formats_yield_the_same_catalogue(
    tmp_path: Path, write: _WriteWorkbook, suffix: str
) -> None:
    """The acceptance criterion「`.xls` と `.xlsx` の**両方**を読める」: the
    same table in either container reads to the same entries."""
    path = write(tmp_path / f"catalogue{suffix}", [_TITLE, _HEADER, _DATA])
    catalog = read_error_catalog(path)
    (entry,) = catalog.entries
    assert (entry.round_label, entry.question_label) == ("第1回", "問1")
    assert (entry.mistake, entry.deduction, entry.red_ink) == ("誤答A", "3点減", "赤入れA")


def test_a_numeric_question_number_does_not_become_a_float(tmp_path: Path) -> None:
    """問題番号 arrives as a number in the measured ``.xlsx``; 「問5.0」 is not
    what the sheet says."""
    path = _write_xlsx(
        tmp_path / "numeric.xlsx", [_TITLE, _HEADER, ["第1回", 5, "", "誤答A", "", "", "赤入れA"]]
    )
    (entry,) = read_error_catalog(path).entries
    assert entry.question_label == "5"


def test_an_unrecognized_layout_raises_instead_of_reading_empty(tmp_path: Path) -> None:
    """The same property as the domain test, but through a real file -- the
    path a caller actually takes."""
    path = _write_xlsx(
        tmp_path / "other.xlsx",
        [["日付", "担当", "メモ"], ["2026-09-10", "架空の氏名", "自由記述"]],
    )
    with pytest.raises(ErrorCatalogUnreadable):
        read_error_catalog(path)


def test_word_is_refused_by_format_not_read_as_empty(tmp_path: Path) -> None:
    """Issue #106 rules Word out of scope. Out of scope must look like a
    refusal, not like a document with nothing in it."""
    path = tmp_path / "material.docx"
    path.write_bytes(b"PK\x03\x04not really a docx")
    with pytest.raises(UnsupportedCatalogFormat):
        read_error_catalog(path)


def _material(
    stored_path: str, role: MaterialRole = MaterialRole.ANNOTATION_RESOURCE
) -> TestMaterial:
    return TestMaterial(
        id=f"material-{stored_path}",
        test_id="test-1",
        role=role,
        stored_path=stored_path,
        sha256="0" * 64,
        size_bytes=1,
        original_filename=None,
        created_at=datetime(2026, 9, 10, tzinfo=UTC),
    )


class TestAnnotationResourceCatalog:
    """The grading-side policy: 添削資料 is optional, so a test without a
    readable one is still graded -- but the three reasons for getting nothing
    must not look alike from the outside."""

    def test_reads_a_registered_excel_resource(self, tmp_path: Path) -> None:
        store = LocalFileStore(tmp_path)
        _write_xlsx(store.resolve_stored_path("tests/test-1/03.xlsx"), [_TITLE, _HEADER, _DATA])
        catalog = annotation_resource_catalog(store, [_material("tests/test-1/03.xlsx")])
        assert catalog is not None
        assert len(catalog.entries) == 1

    def test_no_annotation_resource_registered_is_quiet(
        self, tmp_path: Path, caplog: pytest.LogCaptureFixture
    ) -> None:
        """Nothing registered is the ordinary case for a test whose owner has
        no 添削資料; it is not a fault and must not fill the log."""
        store = LocalFileStore(tmp_path)
        with caplog.at_level(logging.WARNING):
            assert (
                annotation_resource_catalog(
                    store, [_material("tests/test-1/01.pdf", MaterialRole.GRADING_CRITERIA)]
                )
                is None
            )
        assert caplog.records == []

    def test_an_unreadable_resource_is_logged_not_silently_empty(
        self, tmp_path: Path, caplog: pytest.LogCaptureFixture
    ) -> None:
        """The failure this Issue exists to prevent: a registered file whose
        columns are not understood must not be indistinguishable from having
        registered nothing."""
        store = LocalFileStore(tmp_path)
        _write_xlsx(
            store.resolve_stored_path("tests/test-1/03.xlsx"),
            [["日付", "担当"], ["2026-09-10", "架空の氏名"]],
        )
        with caplog.at_level(logging.WARNING):
            assert annotation_resource_catalog(store, [_material("tests/test-1/03.xlsx")]) is None
        assert len(caplog.records) == 1
        message = caplog.records[0].getMessage()
        # "could not be read", not "read 0 entries": the distinction is the
        # whole point of the acceptance criterion.
        assert "no catalogue could be read" in message
        assert "no_header_row" in message

    def test_a_corrupt_file_is_logged_without_quoting_it(
        self, tmp_path: Path, caplog: pytest.LogCaptureFixture
    ) -> None:
        """A library's parse error can quote the bytes it choked on, and those
        bytes are the school's material. Only the exception *type* is logged.
        """
        store = LocalFileStore(tmp_path)
        path = store.resolve_stored_path("tests/test-1/03.xls")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes("架空の秘密文字列".encode())
        with caplog.at_level(logging.WARNING):
            assert annotation_resource_catalog(store, [_material("tests/test-1/03.xls")]) is None
        assert len(caplog.records) == 1
        assert "架空の秘密文字列" not in caplog.records[0].getMessage()

    def test_a_word_resource_is_not_treated_as_an_excel_one(
        self, tmp_path: Path, caplog: pytest.LogCaptureFixture
    ) -> None:
        """The 4 Word-only subjects: no catalogue, and the reader does not
        pretend to have opened the file."""
        store = LocalFileStore(tmp_path)
        path = store.resolve_stored_path("tests/test-1/03.docx")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"PK\x03\x04")
        with caplog.at_level(logging.WARNING):
            assert annotation_resource_catalog(store, [_material("tests/test-1/03.docx")]) is None

    def test_a_partially_readable_workbook_says_so(
        self, tmp_path: Path, caplog: pytest.LogCaptureFixture
    ) -> None:
        """A workbook whose second sheet is a legend still yields a catalogue.
        Reporting only the entries would present a partial read as a whole
        one."""
        store = LocalFileStore(tmp_path)
        path = store.resolve_stored_path("tests/test-1/03.xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        table = workbook.active
        assert table is not None
        for row in (_TITLE, _HEADER, _DATA):
            table.append(row)
        legend = workbook.create_sheet("legend")
        legend.append(["記号", "意味"])
        legend.append(["◎", "架空の説明"])
        workbook.save(path)

        with caplog.at_level(logging.WARNING):
            catalog = annotation_resource_catalog(store, [_material("tests/test-1/03.xlsx")])

        assert catalog is not None
        assert len(catalog.entries) == 1
        assert len(caplog.records) == 1
        assert "sheet 2: no_header_row" in caplog.records[0].getMessage()

    def test_the_oldest_excel_resource_wins(self, tmp_path: Path) -> None:
        """A test may hold several 添削資料; switching between them when
        another is attached would change grading input with nobody asking."""
        store = LocalFileStore(tmp_path)
        _write_xlsx(store.resolve_stored_path("tests/test-1/03_1.xlsx"), [_TITLE, _HEADER, _DATA])
        _write_xlsx(
            store.resolve_stored_path("tests/test-1/03_2.xlsx"),
            [_TITLE, _HEADER, ["第2回", "問9", "", "誤答B", "", "", "赤入れB"]],
        )
        catalog = annotation_resource_catalog(
            store,
            [_material("tests/test-1/03_1.xlsx"), _material("tests/test-1/03_2.xlsx")],
        )
        assert catalog is not None
        assert [entry.question_label for entry in catalog.entries] == ["問1"]
